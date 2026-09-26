"""
tests/test_document_generation.py
==================================
Unit and integration tests for Document Generation Subgraph, Excel/Word generators & validators.
"""

import os
import tempfile
import pytest
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook
from docx import Document

from app.agent.subgraphs.document_generation.generators.excel_generator import generate_excel_document
from app.agent.subgraphs.document_generation.validators.excel_validator import validate_excel_document
from app.agent.subgraphs.document_generation.generators.word_generator import generate_word_document
from app.agent.subgraphs.document_generation.validators.word_validator import validate_word_document
from app.agent.subgraphs.document_generation.graph import document_generation_graph
from app.agent.tools.subgraphs.invoke_document_generation import invoke_document_generation


@pytest.fixture
def temp_dir():
    with tempfile.TemporaryDirectory() as tmpdir:
        yield tmpdir


# ---------------------------------------------------------------------------
# 1. EXCEL TESTS
# ---------------------------------------------------------------------------

def test_excel_generator_and_validator_success(temp_dir):
    spec = {
        "format": "xlsx",
        "title": "Low Stock Inventory Report",
        "filename": "low_stock_report",
        "sheets": [
            {
                "name": "Low Stock",
                "columns": ["Product Name", "Current Stock", "Reorder Level", "Price"],
                "rows": [
                    ["Rice 5kg", 4, 10, 250.0],
                    ["Sugar 1kg", 2, 15, 45.0],
                    ["Oil 1L", 0, 8, 120.0],
                ],
                "currency_columns": ["Price"],
                "include_totals_row": True,
                "freeze_header": True,
                "auto_width": True,
            }
        ],
        "charts": [
            {
                "title": "Stock Levels",
                "chart_type": "bar",
                "sheet_name": "Low Stock",
                "categories_column": "Product Name",
                "values_column": "Current Stock",
            }
        ],
    }

    file_path, filename = generate_excel_document(spec, temp_dir)

    assert os.path.exists(file_path)
    assert filename.endswith(".xlsx")

    # Verify workbook structure with openpyxl
    wb = load_workbook(file_path, data_only=True)
    assert "Low Stock" in wb.sheetnames
    ws = wb["Low Stock"]
    assert ws.cell(row=1, column=1).value == "Low Stock Inventory Report"

    # Validate document deterministically
    val_res = validate_excel_document(file_path, spec)
    assert val_res["valid"] is True
    assert len(val_res["errors"]) == 0


def test_excel_validator_failure_missing_sheet(temp_dir):
    spec = {
        "format": "xlsx",
        "title": "Sample Report",
        "sheets": [{"name": "Actual Sheet", "columns": ["Col1"], "rows": [["Val1"]]}],
    }
    file_path, _ = generate_excel_document(spec, temp_dir)

    invalid_spec = {
        "sheets": [{"name": "NonExistentSheet", "columns": ["Col1"]}]
    }
    val_res = validate_excel_document(file_path, invalid_spec)
    assert val_res["valid"] is False
    assert any("Required sheet 'NonExistentSheet' is missing" in err for err in val_res["errors"])


# ---------------------------------------------------------------------------
# 2. WORD TESTS
# ---------------------------------------------------------------------------

def test_word_generator_and_validator_success(temp_dir):
    spec = {
        "format": "docx",
        "title": "Monthly Business Report",
        "subtitle": "Store #101 Performance Summary",
        "filename": "monthly_report",
        "sections": [
            {
                "heading": "Executive Summary",
                "heading_level": 1,
                "content": "Overall store performance was strong this month with significant revenue growth.",
                "bullet_points": ["Revenue increased by 15%", "Top product: Basmati Rice"],
            },
            {
                "heading": "Sales Breakdown",
                "heading_level": 2,
                "table": {
                    "headers": ["Category", "Monthly Revenue"],
                    "rows": [["Groceries", "₹45,000"], ["Beverages", "₹12,000"]],
                },
            },
        ],
    }

    file_path, filename = generate_word_document(spec, temp_dir)

    assert os.path.exists(file_path)
    assert filename.endswith(".docx")

    doc = Document(file_path)
    text = "\n".join([p.text for p in doc.paragraphs])
    assert "Monthly Business Report" in text
    assert "Executive Summary" in text
    assert len(doc.tables) == 1

    val_res = validate_word_document(file_path, spec)
    assert val_res["valid"] is True
    assert len(val_res["errors"]) == 0


def test_word_validator_failure_missing_heading(temp_dir):
    spec = {
        "format": "docx",
        "title": "Short Note",
        "sections": [{"heading": "Introduction", "content": "Hello"}],
    }
    file_path, _ = generate_word_document(spec, temp_dir)

    invalid_spec = {
        "title": "Short Note",
        "sections": [{"heading": "Missing Financial Analysis", "content": "N/A"}],
    }

    val_res = validate_word_document(file_path, invalid_spec)
    assert val_res["valid"] is False
    assert any("Expected section heading 'Missing Financial Analysis' is missing" in err for err in val_res["errors"])


# ---------------------------------------------------------------------------
# 3. SUBGRAPH INTEGRATION TEST (WITH MOCKED R2)
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_subgraph_execution_success(temp_dir):
    input_state = {
        "user_id": "usr_test123",
        "store_id": "store_test456",
        "document_type": "excel",
        "request": {
            "document_type": "excel",
            "title": "Test Inventory Export",
            "purpose": "Export low stock items",
            "data": [
                {"product": "Wheat Flour", "stock": 5},
                {"product": "Salt 1kg", "stock": 2},
            ],
        },
    }

    mock_r2_service = MagicMock()
    mock_r2_service.generate_storage_key.return_value = "stores/store_test456/users/usr_test123/docs/file_123_test.xlsx"
    mock_r2_service.upload_file.return_value = {
        "bucket": "test-bucket",
        "storage_key": "stores/store_test456/users/usr_test123/docs/file_123_test.xlsx",
        "size_bytes": 1024,
        "mime_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }
    mock_r2_service.generate_presigned_url.return_value = "https://r2.test/download_presigned_url"

    with patch("app.agent.subgraphs.document_generation.nodes.get_r2_storage_service", return_value=mock_r2_service):
        output_state = await document_generation_graph.ainvoke(input_state)

    assert output_state.get("status") == "success"
    file_meta = output_state.get("file_metadata", {})
    assert file_meta.get("document_type") == "excel"
    assert file_meta.get("filename").endswith(".xlsx")
    assert file_meta.get("storage_key") == "stores/store_test456/users/usr_test123/docs/file_123_test.xlsx"
    assert file_meta.get("download_url") == "https://r2.test/download_presigned_url"


@pytest.mark.asyncio
async def test_subgraph_bounded_retry_failure():
    input_state = {
        "user_id": "usr_test",
        "store_id": "store_test",
        "document_type": "excel",
        "max_retries": 1,
        "request": {"title": "Bad Report", "document_type": "excel"},
    }

    # Patch validate_document_node to simulate validation failure
    with patch(
        "app.agent.subgraphs.document_generation.nodes.validate_excel_document",
        return_value={"valid": False, "errors": ["Forced validation error for retry test"], "warnings": []},
    ):
        output_state = await document_generation_graph.ainvoke(input_state)

    assert output_state.get("status") == "failed"
    assert output_state.get("retry_count") >= 1


# ---------------------------------------------------------------------------
# 4. TOOL STUB TEST
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_invoke_document_generation_tool():
    res = await invoke_document_generation.ainvoke({
        "document_type": "excel",
        "title": "Sales Performance",
        "purpose": "Show Q3 Sales",
        "requirements": ["include charts"],
        "data": [{"month": "Jan", "sales": 100}],
    })

    assert res["__subgraph__"] == "document_generation"
    assert res["request"]["title"] == "Sales Performance"
    assert res["request"]["document_type"] == "excel"
