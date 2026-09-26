"""
app/agent/subgraphs/document_generation/validators/word_validator.py
=====================================================================
Deterministic validation for generated Word (.docx) files.
"""

import os
from typing import Any
import structlog
from docx import Document

from app.config.settings import get_settings

LOGGER = structlog.get_logger("vyaparsathi.ai.doc_gen.validators.word")


def validate_word_document(file_path: str, spec: dict[str, Any]) -> dict[str, Any]:
    """
    Perform deterministic validation on a generated Word file against its specification.

    Checks:
    - File existence & size
    - Corrupt file / python-docx loading check
    - Document title present
    - Headings present
    - Required tables present

    Returns:
        dict with keys: "valid" (bool), "errors" (list[str]), "warnings" (list[str])
    """
    errors: list[str] = []
    warnings: list[str] = []

    # 1. File existence
    if not os.path.exists(file_path):
        return {
            "valid": False,
            "errors": [f"File does not exist at path: {file_path}"],
            "warnings": [],
        }

    # 2. File size check
    settings = get_settings()
    max_bytes = settings.max_generated_file_size_mb * 1024 * 1024
    file_size = os.path.getsize(file_path)

    if file_size == 0:
        return {
            "valid": False,
            "errors": ["Generated Word file is 0 bytes (empty file)."],
            "warnings": [],
        }

    if file_size > max_bytes:
        return {
            "valid": False,
            "errors": [f"Generated Word file size ({file_size} bytes) exceeds limit ({max_bytes} bytes)."],
            "warnings": [],
        }

    # 3. python-docx load check (Corruption test)
    try:
        doc = Document(file_path)
    except Exception as exc:
        LOGGER.error("word_validator.load_corrupt", path=file_path, error=str(exc))
        return {
            "valid": False,
            "errors": [f"Word document is corrupt or unreadable by python-docx: {str(exc)}"],
            "warnings": [],
        }

    # 4. Check Title
    expected_title = spec.get("title", "")
    all_text = "\n".join([p.text for p in doc.paragraphs])

    if expected_title and expected_title not in all_text:
        errors.append(f"Expected document title '{expected_title}' was not found in paragraphs.")

    # 5. Check Headings & Tables in Sections
    sections_spec = spec.get("sections", [])
    headings_in_doc = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]

    expected_tables_count = 0
    for sec in sections_spec:
        h_text = sec.get("heading")
        if h_text and h_text not in headings_in_doc and h_text not in all_text:
            errors.append(f"Expected section heading '{h_text}' is missing from document.")
        
        if sec.get("table"):
            expected_tables_count += 1

    actual_tables_count = len(doc.tables)
    if expected_tables_count > 0 and actual_tables_count < expected_tables_count:
        errors.append(
            f"Expected {expected_tables_count} table(s) in document, but only found {actual_tables_count}."
        )

    is_valid = len(errors) == 0

    LOGGER.info(
        "word_validator.completed",
        valid=is_valid,
        error_count=len(errors),
        warning_count=len(warnings),
    )

    return {
        "valid": is_valid,
        "errors": errors,
        "warnings": warnings,
    }
