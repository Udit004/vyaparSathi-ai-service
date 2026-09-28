"""
app/agent/subgraphs/document_generation/validators/excel_validator.py
=======================================================================
Deterministic and semantic validation for generated Excel (.xlsx) files.
"""

import os
from typing import Any
import structlog
from openpyxl import load_workbook

from app.config.settings import get_settings

LOGGER = structlog.get_logger("vyaparsathi.ai.doc_gen.validators.excel")


def validate_excel_document(file_path: str, spec: dict[str, Any]) -> dict[str, Any]:
    """
    Perform deterministic validation on a generated Excel file against its specification.

    Checks:
    - File existence & size
    - Corrupt file / openpyxl loading check
    - Sheet names present
    - Column headers present
    - Non-empty row count
    - Charts presence if specified in spec

    Returns:
        dict with keys: "valid" (bool), "errors" (list[str]), "warnings" (list[str])
    """
    errors: list[str] = []
    warnings: list[str] = []

    # 1. File existence
    if not os.path.exists(file_path):
        return {
            "valid": False,
            "errors": [f"File does not exist at path: {file_path}"],
            "warnings": [],
        }

    # 2. File size check
    settings = get_settings()
    max_bytes = settings.max_generated_file_size_mb * 1024 * 1024
    file_size = os.path.getsize(file_path)

    if file_size == 0:
        return {
            "valid": False,
            "errors": ["Generated file is 0 bytes (empty file)."],
            "warnings": [],
        }

    if file_size > max_bytes:
        return {
            "valid": False,
            "errors": [f"Generated file size ({file_size} bytes) exceeds limit ({max_bytes} bytes)."],
            "warnings": [],
        }

    # 3. openpyxl load check (Corruption test)
    try:
        wb = load_workbook(file_path, data_only=True)
    except Exception as exc:
        LOGGER.error("excel_validator.load_corrupt", path=file_path, error=str(exc))
        return {
            "valid": False,
            "errors": [f"Workbook is corrupt or unreadable by openpyxl: {str(exc)}"],
            "warnings": [],
        }

    # 4. Sheet check
    expected_sheets = spec.get("sheets", [])
    sheet_names_in_wb = wb.sheetnames

    for sheet_spec in expected_sheets:
        expected_name = sheet_spec.get("name", "")[:31]  # 31 char limit
        if expected_name and expected_name not in sheet_names_in_wb:
            errors.append(f"Required sheet '{expected_name}' is missing in generated workbook.")
            continue

        ws = wb[expected_name]

        # 5. Column check
        expected_cols = sheet_spec.get("columns", [])
        if expected_cols:
            # Find header row (row 3 in default excel_generator)
            actual_cols = []
            for row in range(1, 10):
                row_vals = [str(cell.value) for cell in ws[row] if cell.value is not None]
                if any(col in row_vals for col in expected_cols):
                    actual_cols = row_vals
                    break
            
            for col in expected_cols:
                if col not in actual_cols:
                    errors.append(f"Sheet '{expected_name}' is missing expected column: '{col}'.")

        # 6. Row count check
        expected_rows = sheet_spec.get("rows", [])
        if expected_rows and ws.max_row <= 3:
            errors.append(f"Sheet '{expected_name}' contains no data rows (max row is {ws.max_row}).")

    # 7. Chart check
    expected_charts = spec.get("charts", [])
    if expected_charts:
        total_charts_found = 0
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            total_charts_found += len(ws._charts)
        
        if total_charts_found < len(expected_charts):
            warnings.append(
                f"Expected {len(expected_charts)} chart(s), but found {total_charts_found} chart(s) in workbook."
            )

    is_valid = len(errors) == 0

    LOGGER.info(
        "excel_validator.completed",
        valid=is_valid,
        error_count=len(errors),
        warning_count=len(warnings),
    )

    return {
        "valid": is_valid,
        "errors": errors,
        "warnings": warnings,
    }
