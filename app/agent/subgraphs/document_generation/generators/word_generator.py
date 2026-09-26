"""
app/agent/subgraphs/document_generation/generators/word_generator.py
======================================================================
Deterministic Word document generator using python-docx.
Translates structured WordDocumentSpec into a valid .docx file.
"""

import os
import uuid
from typing import Any
import structlog
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

LOGGER = structlog.get_logger("vyaparsathi.ai.doc_gen.generators.word")


def _set_cell_background(cell, hex_color: str):
    """Set background color of a table cell in docx."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def generate_word_document(spec: dict[str, Any], output_dir: str) -> tuple[str, str]:
    """
    Generate a Word document (.docx) deterministically using python-docx based on spec.

    Args:
        spec: Structured Word specification dict (matching WordDocumentSpec schema)
        output_dir: Directory where temporary .docx file will be created

    Returns:
        tuple[file_path, filename]
    """
    os.makedirs(output_dir, exist_ok=True)

    file_id = f"doc_{uuid.uuid4().hex[:12]}"
    raw_filename = spec.get("filename") or spec.get("title") or "document"
    safe_basename = "".join(c if c.isalnum() or c in ("-", "_") else "_" for c in raw_filename.lower())
    safe_basename = safe_basename.strip("_") or "document"
    filename = f"{safe_basename}_{file_id[:8]}.docx"
    file_path = os.path.join(output_dir, filename)

    doc = Document()

    # Set page margins (1 inch standard)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Document Title
    title_text = spec.get("title", "Document Report")
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title_p.add_run(title_text)
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(31, 78, 121)  # Navy blue

    # Subtitle / Metadata header
    subtitle_text = spec.get("subtitle")
    if subtitle_text:
        sub_p = doc.add_paragraph()
        sub_run = sub_p.add_run(subtitle_text)
        sub_run.font.name = "Calibri"
        sub_run.font.size = Pt(12)
        sub_run.font.italic = True
        sub_run.font.color.rgb = RGBColor(89, 89, 89)

    doc.add_paragraph()  # spacing

    # Sections processing
    sections_spec = spec.get("sections", [])
    for sec in sections_spec:
        heading_text = sec.get("heading")
        if heading_text:
            level = sec.get("heading_level", 1)
            doc.add_heading(heading_text, level=min(max(level, 1), 4))

        # Main paragraph content
        content_text = sec.get("content")
        if content_text:
            p = doc.add_paragraph(content_text)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(6)

        # Bullet points
        bullet_points = sec.get("bullet_points", [])
        for bp in bullet_points:
            doc.add_paragraph(bp, style="List Bullet")

        # Numbered points
        numbered_points = sec.get("numbered_points", [])
        for np in numbered_points:
            doc.add_paragraph(np, style="List Number")

        # Table embedded in section
        table_spec = sec.get("table")
        if table_spec:
            headers = table_spec.get("headers", [])
            rows = table_spec.get("rows", [])

            if headers:
                table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = True

                # Format Header Row
                hdr_cells = table.rows[0].cells
                for i, header_text in enumerate(headers):
                    hdr_cells[i].text = str(header_text)
                    _set_cell_background(hdr_cells[i], "1F4E79")
                    for p in hdr_cells[i].paragraphs:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        for run in p.runs:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)

                # Format Data Rows
                for r_idx, row_data in enumerate(rows):
                    row_cells = table.rows[r_idx + 1].cells
                    bg_color = "F2F2F2" if r_idx % 2 == 1 else "FFFFFF"
                    for c_idx, val in enumerate(row_data):
                        if c_idx < len(row_cells):
                            row_cells[c_idx].text = str(val) if val is not None else ""
                            _set_cell_background(row_cells[c_idx], bg_color)

                doc.add_paragraph()  # spacing after table

    doc.save(file_path)
    LOGGER.info("word_generator.created", path=file_path, filename=filename)
    return file_path, filename
